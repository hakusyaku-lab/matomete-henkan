import json

import pytest
from PIL import Image
from pypdf import PdfReader, PdfWriter

from converters import csv_tools, image_tools, mapping, normalize, pdf_tools, text_io


# ---------- 文字コード判定 ----------

@pytest.mark.parametrize("enc,expected", [("cp932", "cp932"), ("utf-8", "utf-8"), ("utf-8-sig", "utf-8-sig")])
def test_read_table_detects_encoding(tmp_path, enc, expected):
    p = tmp_path / "a.csv"
    p.write_bytes("氏名,電話番号\n山田 太郎,090-1234-5678\n㈱テスト,0312345678\n".encode(enc))
    t = text_io.read_table(p)
    assert t["encoding"] == expected
    assert t["rows"][1] == ["山田 太郎", "090-1234-5678"]
    assert t["row_count"] == 2


def test_read_table_tab_delimited(tmp_path):
    p = tmp_path / "a.txt"
    p.write_text("a\tb\n1\t2\n", encoding="utf-8")
    assert text_io.read_table(p)["rows"] == [["a", "b"], ["1", "2"]]


def test_find_unencodable():
    found = text_io.find_unencodable([["名前"], ["絵文字😀"], ["①㈱"]])
    assert [f["char"] for f in found] == ["😀"]  # ①㈱ は Shift_JIS(cp932) で表せる


# ---------- 整形 ----------

def test_normalize_functions():
    assert normalize.to_hankaku_alnum("ＡＢＣ１２３－あ") == "ABC123－あ"
    assert normalize.hankaku_kana_to_zenkaku("ｶﾞｲﾄﾞ abc") == "ガイド abc"
    assert normalize.wareki_to_seireki("令和6年1月2日") == "2024/01/02"
    assert normalize.wareki_to_seireki("平成元年12月31日 生") == "1989/12/31 生"
    assert normalize.wareki_to_seireki("R6.1.2") == "2024/01/02"
    assert normalize.wareki_to_seireki("R6-1-2X") == "R6-1-2X"
    assert normalize.wareki_to_seireki("令和6年13月2日") == "令和6年13月2日"
    assert normalize.replace_kishu("㈱山田①㎏") == "(株)山田(1)kg"
    assert normalize.remove_comma("1,280") == "1280"
    assert normalize.remove_comma("東京,大阪") == "東京,大阪"


def test_normalize_rows_keeps_header():
    rows = [["ＩＤ"], ["００１"]]
    assert normalize.normalize_rows(rows, {"hankaku": True}) == [["ＩＤ"], ["001"]]


# ---------- 結合と書き出し ----------

def test_merge_tables_aligns_columns():
    t1 = {"name": "a.csv", "rows": [["コード", "名前"], ["001", "りんご"]]}
    t2 = {"name": "b.csv", "rows": [["名前", "価格", "コード"], ["みかん", "100", "002"]]}
    merged = csv_tools.merge_tables([t1, t2], add_source=True)
    assert merged[0] == ["元ファイル", "コード", "名前", "価格"]
    assert merged[1] == ["a.csv", "001", "りんご", ""]
    assert merged[2] == ["b.csv", "002", "みかん", "100"]


def test_xlsx_protects_leading_zero_and_formula(tmp_path):
    from openpyxl import load_workbook

    rows = [["電話", "金額", "メモ"], ["0901234", "1200", "=1+1"]]
    p = tmp_path / "o.xlsx"
    csv_tools.write_xlsx(p, rows, protect_zero=True)
    ws = load_workbook(p).active
    assert ws["A2"].value == "0901234"
    assert ws["B2"].value == "1200"
    assert ws["C2"].value == "=1+1" and ws["C2"].data_type == "s"

    csv_tools.write_xlsx(p, rows, protect_zero=False)
    ws = load_workbook(p).active
    assert ws["A2"].value == "0901234"  # 先頭0は数値にしない
    assert ws["B2"].value == 1200


def test_export_csv_sjis_and_json(tmp_path):
    rows = [["名前", "メモ"], ["山田", "😀あ"]]
    p = tmp_path / "o.csv"
    csv_tools.export_rows(p, rows, "csv_sjis")
    assert p.read_bytes().decode("cp932").splitlines()[1] == "山田,?あ"
    j = tmp_path / "o.json"
    csv_tools.export_rows(j, rows, "json")
    assert json.loads(j.read_text(encoding="utf-8")) == [{"名前": "山田", "メモ": "😀あ"}]


# ---------- 画像 ----------

def _rgba_png(path, size=(200, 100)):
    img = Image.new("RGBA", size, (255, 0, 0, 0))
    img.paste((0, 0, 255, 255), (50, 25, 150, 75))
    img.save(path)


def test_image_jpeg_square_resize(tmp_path):
    src = tmp_path / "in.png"
    _rgba_png(src)
    info = image_tools.inspect_image(src)
    assert info["has_alpha"] and info["thumb"].startswith("data:image/jpeg")
    res = image_tools.convert_image(src, tmp_path / "out", {"fmt": "jpeg", "max_side": 100, "square": True, "quality": 85})
    out = Image.open(res["output"])
    assert out.format == "JPEG" and out.size == (100, 100)
    assert out.getpixel((2, 2)) == (255, 255, 255) or sum(out.getpixel((2, 2))) > 740  # 余白・透明部分は白
    assert "透明部分を白にしました" not in res["notes"]  # 正方形化で先に白にしている


def test_image_png_keeps_alpha_and_sequential(tmp_path):
    src = tmp_path / "in.png"
    _rgba_png(src)
    res = image_tools.convert_image(src, tmp_path, {"fmt": "png", "sequential": True, "prefix": "item"}, index=3)
    assert res["name"] == "item_0003.png"
    assert Image.open(res["output"]).mode == "RGBA"
    res2 = image_tools.convert_image(src, tmp_path, {"fmt": "webp", "lossless": True, "sequential": True, "prefix": "item"}, index=3)
    assert res2["name"] == "item_0003.webp"


def test_image_strips_gps(tmp_path):
    src = tmp_path / "gps.jpg"
    img = Image.new("RGB", (40, 30), (10, 200, 10))
    exif = img.getexif()
    exif[0x010F] = "TestMaker"
    exif[0x8825] = {1: "N", 2: (35.0, 0.0, 0.0)}
    img.save(src, exif=exif)
    assert 0x8825 in Image.open(src).getexif()
    res = image_tools.convert_image(src, tmp_path / "o", {"fmt": "jpeg", "strip_gps": True})
    e = Image.open(res["output"]).getexif()
    assert 0x8825 not in e and e.get(0x010F) == "TestMaker"


# ---------- 列の対応付け ----------

def test_apply_mapping_and_templates(tmp_path, monkeypatch):
    monkeypatch.setenv("MATOMETE_CONFIG_DIR", str(tmp_path / "cfg"))
    rows = [["商品番号", "販売価格", "商品名"], ["A01", "1,280", "ＴシャツＬ"]]
    m = [
        {"target": "コード", "source": "商品番号", "ops": []},
        {"target": "価格", "source": "販売価格", "ops": ["remove_comma"]},
        {"target": "名前", "source": "商品名", "ops": ["hankaku"]},
        {"target": "送料", "source": "__fixed__", "fixed": "送料無料", "ops": []},
        {"target": "空", "source": "", "ops": []},
    ]
    assert mapping.apply_mapping(rows, m) == [["コード", "価格", "名前", "送料", "空"], ["A01", "1280", "TシャツL", "送料無料", ""]]
    mapping.save_template("モールA→B", {"mapping": m})
    assert mapping.list_templates() == ["モールA→B"]
    assert mapping.load_template("モールA→B")["mapping"] == m
    mapping.delete_template("モールA→B")
    assert mapping.list_templates() == []
    auto = mapping.auto_mapping(["商品名", "価格"], ["価格", "説明"])
    assert [a["source"] for a in auto] == ["価格", ""]


# ---------- PDF ----------

def _pdf(path, n):
    w = PdfWriter()
    for i in range(n):
        w.add_blank_page(width=200 + i, height=300)
    with open(path, "wb") as f:
        w.write(f)


def test_parse_ranges():
    assert pdf_tools.parse_ranges("1-3, 5", 6) == [0, 1, 2, 4]
    assert pdf_tools.parse_ranges("５、２－", 6) == [4, 1, 2, 3, 4, 5]
    with pytest.raises(ValueError):
        pdf_tools.parse_ranges("7", 6)
    with pytest.raises(ValueError):
        pdf_tools.parse_ranges("abc", 6)


def test_pdf_merge_split_extract(tmp_path):
    a, b = tmp_path / "a.pdf", tmp_path / "b.pdf"
    _pdf(a, 3)
    _pdf(b, 2)
    info = pdf_tools.inspect_pdf(a)
    assert info["pages"] == 3 and len(info["thumbs"]) == 3
    m = tmp_path / "m.pdf"
    assert pdf_tools.merge_pdfs([str(a), str(b)], m) == 5
    assert len(PdfReader(m).pages) == 5
    outs = pdf_tools.split_pdf(a, tmp_path / "s", every=2)
    assert [p.split("/")[-1].split("\\")[-1] for p in outs] == ["a_p001-002.pdf", "a_p003.pdf"]
    e = tmp_path / "e.pdf"
    assert pdf_tools.extract_pages(a, "1,3", e) == 2
    assert [round(float(p.mediabox.width)) for p in PdfReader(e).pages] == [200, 202]


# ---------- アプリ本体（ダイアログの拡張子フィルター） ----------

def test_dialog_filters_are_valid():
    pytest.importorskip("webview")
    from webview.util import parse_file_type

    import main

    for filters in main.FILE_TYPES.values():
        for f in filters:
            parse_file_type(f)
    for ext, _ in csv_tools.FORMATS.values():
        parse_file_type(f"{ext.lstrip('.').upper()} (*{ext})")
